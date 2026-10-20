import io
import os
import sys
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import load_workbook

import PurviewClassificationsReport as report


class SourceMatchingTests(unittest.TestCase):
    def setUp(self):
        self.source = {
            "name": "awss3-iodemodata",
            "kind": "AmazonS3",
            "properties": {"serviceUrl": "s3://iodemodata"},
        }
        self.other_source = {
            "name": "awss3-iopurviewdemodata",
            "kind": "AmazonS3",
            "properties": {"serviceUrl": "s3://iopurviewdemodata"},
        }
        self.qualified_name = "s3://iodemodata/SF-Test/purview-sample-schema.csv"

    def test_s3_service_url_produces_bucket_locator(self):
        for service_url in ("s3://iodemodata", " s3://iodemodata/ "):
            with self.subTest(service_url=service_url):
                self.source["properties"]["serviceUrl"] = service_url
                self.assertEqual(report.data_source_locators(self.source), ["iodemodata"])

    def test_nested_service_url_is_recognized(self):
        self.source["properties"] = {
            "connection": {"serviceUrl": "s3://iodemodata"}
        }
        self.assertEqual(report.data_source_locators(self.source), ["iodemodata"])

    def test_s3_files_and_columns_match_in_selected_and_all_modes(self):
        for selected_name in (None, self.source["name"]):
            index = report.build_source_index(
                [self.other_source, self.source], selected_name
            )
            for qualified_name in (
                "s3://iodemodata",
                self.qualified_name,
                f"{self.qualified_name}/routing_number",
            ):
                with self.subTest(
                    selected_name=selected_name, qualified_name=qualified_name
                ):
                    self.assertIs(
                        report.match_asset_to_source(
                            {"qualifiedName": qualified_name}, index
                        ),
                        self.source,
                    )
            self.assertIs(
                report.match_asset_to_source(
                    {"qualifiedName": "s3://iopurviewdemodata/other.csv"}, index
                ),
                self.other_source,
            )
            self.assertIsNone(
                report.match_asset_to_source(
                    {"qualifiedName": "s3://iodemodata-backup/other.csv"}, index
                )
            )

    def test_existing_endpoint_fields_still_match(self):
        for key in (
            "endpoint", "host", "hostName", "qualifiedName",
            "resourceUri", "serverEndpoint", "url",
        ):
            with self.subTest(key=key):
                source = {
                    "name": "existing-source",
                    "properties": {key: "https://storage.example.com"},
                }
                self.assertIs(
                    report.match_asset_to_source(
                        {"qualifiedName": "https://storage.example.com/data/file.csv"},
                        report.build_source_index([source]),
                    ),
                    source,
                )

    def test_source_without_endpoint_still_requires_explicit_prefix(self):
        self.source["properties"] = {"roleARN": "arn:aws:iam::123456789012:role/scan"}
        self.assertEqual(report.data_source_locators(self.source), [])
        index = report.build_source_index(
            [self.source], self.source["name"], "s3://iodemodata"
        )
        self.assertIs(
            report.match_asset_to_source({"qualifiedName": self.qualified_name}, index),
            self.source,
        )

    def test_more_specific_registration_wins(self):
        folder_source = {
            "name": "s3-folder",
            "properties": {"serviceUrl": "s3://iodemodata/SF-Test"},
        }
        index = report.build_source_index([self.source, folder_source])
        self.assertIs(
            report.match_asset_to_source({"qualifiedName": self.qualified_name}, index),
            folder_source,
        )

    def test_connector_specific_registration_locators(self):
        cases = [
            ("AzureCosmosDb", {"accountUri": "https://cosmos.documents.azure.com:443/"},
             "https://cosmos.documents.azure.com/dbs/demo/colls/users"),
            ("Dataverse", {"webApiEndpoint": "https://org.crm.dynamics.com/api/data/v9.2"},
             "dataverse://org.crm.dynamics.com/tables/contact"),
            ("DatabricksHms", {"workspaceUrl": "adb-123.azuredatabricks.net"},
             "databricks://adb-123.azuredatabricks.net"),
            ("DatabricksUnityCatalog", {"metastoreId": "test-metastore-id"},
             "databricks://test-metastore-id/catalogs/demo/schemas/public/tables/users"),
            ("MongoDB", {"server": "mongo.example.com:27017"},
             "mongodb://mongo.example.com:27017/databases/demo/collections/users"),
            ("AzureSynapseWorkspace", {"serverlessSqlEndpoint": "demo-ondemand.sql.azuresynapse.net"},
             "mssql://demo-ondemand.sql.azuresynapse.net/serverlessdb"),
            ("AzureSynapseWorkspace", {"dedicatedSqlEndpoint": "demo.sql.azuresynapse.net"},
             "https://demo.azuresynapse.net"),
            ("AzureSynapseWorkspace", {"resourceId": "/subscriptions/123/providers/Microsoft.Synapse/workspaces/demo"},
             "/subscriptions/123/providers/Microsoft.Synapse/workspaces/demo/pipelines/copy"),
            ("AzureSqlDatabaseManagedInstance", {"serverEndpoint": "mi.public.example.net,3342"},
             "mssql://mi.public.example.net/demo/dbo/users"),
            ("AdlsGen2", {"endpoint": "https://storage.dfs.core.windows.net/"},
             "https://storage.core.windows.net"),
            ("Snowflake", {"host": "primary.snowflakecomputing.com",
                           "additionalHosts": ["secondary.snowflakecomputing.com"]},
             "snowflake://secondary.snowflakecomputing.com/demo/public/users"),
        ]
        for kind, properties, qualified_name in cases:
            with self.subTest(kind=kind, qualified_name=qualified_name):
                source = {"name": "registered-source", "kind": kind, "properties": properties}
                index = report.build_source_index([source])
                self.assertIs(
                    report.match_asset_to_source({"qualifiedName": qualified_name}, index),
                    source,
                )

    def test_port_normalization_preserves_nondefault_http_ports(self):
        self.assertEqual(
            report.normalize_locator("https://example.com:443/data"), "example.com/data"
        )
        self.assertEqual(
            report.normalize_locator("http://example.com:80/data"), "example.com/data"
        )
        self.assertEqual(
            report.normalize_locator("https://example.com:8443/data"), "example.com:8443/data"
        )
        self.assertFalse(report.locator_matches(
            "https://example.com:8443/data", "https://example.com:443/"
        ))

    def test_unknown_source_is_not_assigned_by_collection_or_registration_name(self):
        self.source["properties"]["collection"] = {"referenceName": "shared-collection"}
        unknown = {
            "qualifiedName": "s3://unregistered-bucket/awss3-iodemodata/file.csv",
            "collectionId": "shared-collection",
        }
        grouped, summaries, unmatched = report.aggregate_assets(
            [self.source], [unknown], report.build_source_index([self.source])
        )
        self.assertEqual(unmatched, [unknown])
        self.assertEqual(grouped[self.source["name"]], [])
        self.assertEqual(summaries[self.source["name"]]["matched_assets"], 0)

    def test_s3_command_writes_endpoint_assets_and_column_classifications(self):
        asset = {
            "id": "file-guid",
            "name": "purview-sample-schema.csv",
            "qualifiedName": self.qualified_name,
            "entityType": "aws_s3_object",
        }
        column = {
            "id": "column-guid",
            "name": "routing_number",
            "qualifiedName": f"{self.qualified_name}/routing_number",
            "entityType": "column",
            "classification": ["MICROSOFT.FINANCIAL.US.ABA_ROUTING_NUMBER"],
            "parentGuid": asset["id"],
            "parentName": asset["name"],
        }
        credential = Mock()
        credential.get_token.return_value.token = "test-token"
        stdout, stderr = io.StringIO(), io.StringIO()

        with TemporaryDirectory() as directory:
            argv = [
                "PurviewClassificationsReport.py",
                "--data-source", self.source["name"],
                "--output-directory", directory,
            ]
            with (
                patch.object(sys, "argv", argv),
                patch.dict(os.environ, {"PURVIEW_ACCOUNT_NAME": "test-account"}),
                patch.object(report, "load_dotenv"),
                patch.object(report, "create_credential", return_value=credential),
                patch.object(
                    report, "list_data_sources",
                    return_value=[self.other_source, self.source],
                ),
                patch.object(
                    report, "list_catalog_assets",
                    return_value=[
                        asset,
                        {"id": "other-guid", "qualifiedName": "s3://iopurviewdemodata/a.csv"},
                    ],
                ),
                patch.object(
                    report, "list_classified_columns", return_value=[column]
                ) as list_columns,
                redirect_stdout(stdout),
                redirect_stderr(stderr),
            ):
                self.assertEqual(report.main(), 0, stderr.getvalue())

            self.assertEqual(list_columns.call_args.args[3], [asset])
            workbook = load_workbook(Path(directory) / "classifications.xlsx")
            try:
                source_rows = list(workbook["Data Sources"].values)
                source_values = dict(zip(source_rows[0], source_rows[1]))
                self.assertEqual(len(source_rows), 2)
                self.assertEqual(source_values["Data Source"], self.source["name"])
                self.assertEqual(source_values["Endpoint"], "s3://iodemodata")
                self.assertEqual(source_values["Asset Match Locators"], "iodemodata")
                self.assertEqual(source_values["Matched Assets"], 1)
                column_rows = list(workbook["Column Classifications"].values)
                self.assertEqual(len(column_rows), 2)
                column_values = dict(zip(column_rows[0], column_rows[1]))
                self.assertEqual(column_values["Data Source"], self.source["name"])
                self.assertEqual(column_values["Column GUID"], column["id"])
                self.assertEqual(
                    column_values["Classification"], column["classification"][0]
                )
            finally:
                workbook.close()

    def test_all_report_preserves_unmatched_records_without_warning(self):
        unknown = {
            "id": "unregistered-guid", "name": "Legacy table",
            "entityType": "postgresql_table", "assetType": ["PostgreSQL"],
            "qualifiedName": "postgresql://old.example.com/demo/table",
            "classification": ["Sensitive"], "collectionId": "shared-collection",
        }
        credential = Mock()
        credential.get_token.return_value.token = "test-token"
        for per_source in (False, True):
            with self.subTest(per_source=per_source), TemporaryDirectory() as directory:
                argv = [
                    "PurviewClassificationsReport.py", "--data-source", "ALL",
                    "--filename-prefix", "FULL", "--output-directory", directory,
                ]
                if per_source:
                    argv.append("--file-per-data-source")
                stdout, stderr = io.StringIO(), io.StringIO()
                with (
                    patch.object(sys, "argv", argv),
                    patch.dict(os.environ, {"PURVIEW_ACCOUNT_NAME": "test-account"}),
                    patch.object(report, "load_dotenv"),
                    patch.object(report, "create_credential", return_value=credential),
                    patch.object(report, "list_data_sources", return_value=[self.source]),
                    patch.object(report, "list_catalog_assets", return_value=[unknown]),
                    patch.object(report, "list_classified_columns", return_value=[]),
                    redirect_stdout(stdout), redirect_stderr(stderr),
                ):
                    self.assertEqual(report.main(), 0, stderr.getvalue())
                self.assertEqual(stderr.getvalue(), "")
                self.assertIn("Coverage: 1 catalog assets", stdout.getvalue())
                output_files = list(Path(directory).glob("*.xlsx"))
                self.assertEqual(len(output_files), 1)
                expected_name = f"{self.source['name']}-FULL.xlsx" if per_source else "FULL.xlsx"
                self.assertEqual(output_files[0].name, expected_name)
                workbook = load_workbook(output_files[0])
                try:
                    rows = list(workbook["Unmatched Assets"].values)
                    self.assertEqual(len(rows), 2)
                    record = dict(zip(rows[0], rows[1]))
                    self.assertEqual(record["Asset GUID"], unknown["id"])
                    self.assertEqual(record["Classification"], "Sensitive")
                    self.assertEqual(record["Qualified Name"], unknown["qualifiedName"])
                finally:
                    workbook.close()


class FabricWorkspaceMatchingTests(unittest.TestCase):
    def setUp(self):
        self.source = {
            "name": "Fabric", "kind": "Fabric",
            "properties": {"host": "api.powerbi.com", "tenant": "tenant-one"},
        }
        self.workspace = {
            "id": "workspace-guid", "entityType": "fabric_workspace",
            "qualifiedName": "https://app.powerbi.com/groups/workspace-one",
        }
        self.entity = {
            "guid": "workspace-guid", "typeName": "fabric_workspace",
            "attributes": {
                "qualifiedName": self.workspace["qualifiedName"], "tenantId": "tenant-one",
            },
        }

    def resolve(self, sources=None, assets=None):
        return report.add_fabric_workspace_locators(
            Mock(), "https://test.purview.azure.com", {},
            [self.workspace] if assets is None else assets,
            report.build_source_index(sources or [self.source]),
        )

    def test_verified_workspace_matches_both_app_hosts_not_other_workspaces(self):
        with patch.object(report, "request_json", return_value={"entities": [self.entity]}):
            index = self.resolve()
        for host in ("app.fabric.microsoft.com", "app.powerbi.com"):
            self.assertIs(report.match_asset_to_source({
                "qualifiedName": f"https://{host}/groups/workspace-one/lakehouses/lake/tables/t"
            }, index), self.source)
            self.assertIsNone(report.match_asset_to_source({
                "qualifiedName": f"https://{host}/groups/workspace-two/lakehouses/lake"
            }, index))

    def test_wrong_missing_or_ambiguous_tenant_is_not_matched(self):
        for tenant in ("other-tenant", None, "tenant-one"):
            with self.subTest(tenant=tenant):
                self.entity["attributes"]["tenantId"] = tenant
                sources = [self.source]
                if tenant == "tenant-one":
                    sources.append({
                        "name": "Duplicate", "kind": "Fabric",
                        "properties": {"tenant": "tenant-one"},
                    })
                with patch.object(report, "request_json", return_value={"entities": [self.entity]}):
                    index = self.resolve(sources=sources)
                self.assertIsNone(report.match_asset_to_source({
                    "qualifiedName": self.workspace["qualifiedName"] + "/datasets/data",
                }, index))

    def test_shared_public_host_is_not_a_registration_locator(self):
        self.assertEqual(report.data_source_locators(self.source), [])

    def test_missing_workspace_is_resolved_without_modified_time_filter(self):
        asset = {
            "id": "table", "entityType": "fabric_lakehouse_table",
            "qualifiedName": self.workspace["qualifiedName"] + "/lakehouses/lake/tables/t",
        }
        with patch.object(report, "request_json", side_effect=[
            {"value": [], "continuationToken": "next-page"},
            {"value": [self.workspace]},
            {"entities": [self.entity]},
        ]) as request:
            index = self.resolve(assets=[asset])
        self.assertIs(report.match_asset_to_source(asset, index), self.source)
        query = request.call_args_list[0].kwargs["json"]
        self.assertNotIn("timerange", str(query))
        self.assertEqual(request.call_args_list[1].kwargs["json"]["continuationToken"], "next-page")

    def test_missing_workspace_details_remain_an_error(self):
        with patch.object(report, "request_json", return_value={"entities": []}):
            with self.assertRaisesRegex(report.PurviewApiError, "requested workspace"):
                self.resolve()

    def test_tenant_only_registration_can_generate_selected_report(self):
        self.source["properties"].pop("host")
        credential = Mock()
        credential.get_token.return_value.token = "test-token"
        stdout, stderr = io.StringIO(), io.StringIO()
        with TemporaryDirectory() as directory:
            with (
                patch.object(sys, "argv", [
                    "PurviewClassificationsReport.py", "--data-source", "Fabric",
                    "--output-directory", directory,
                ]),
                patch.dict(os.environ, {"PURVIEW_ACCOUNT_NAME": "test-account"}),
                patch.object(report, "load_dotenv"),
                patch.object(report, "create_credential", return_value=credential),
                patch.object(report, "list_data_sources", return_value=[self.source]),
                patch.object(report, "list_catalog_assets", return_value=[self.workspace]),
                patch.object(report, "request_json", return_value={"entities": [self.entity]}),
                patch.object(report, "list_classified_columns", return_value=[]),
                redirect_stdout(stdout), redirect_stderr(stderr),
            ):
                self.assertEqual(report.main(), 0, stderr.getvalue())
            self.assertTrue((Path(directory) / "classifications.xlsx").is_file())
            self.assertEqual(stderr.getvalue(), "")

    def test_non_fabric_registration_does_not_request_workspace_details(self):
        with patch.object(report, "request_json") as request:
            self.resolve(sources=[{"name": "s3", "kind": "AmazonS3", "properties": {}}])
        request.assert_not_called()


if __name__ == "__main__":
    unittest.main()
