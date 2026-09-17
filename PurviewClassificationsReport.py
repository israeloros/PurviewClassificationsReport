"""Generate Excel reports of Microsoft Purview asset classifications.

The script authenticates with either
:class:`azure.identity.DefaultAzureCredential` or
:class:`azure.identity.ClientSecretCredential`, retrieves registered data
sources and catalog assets from the Microsoft Purview data-plane APIs, maps
assets to registrations by normalized locator, and writes one combined
workbook or one workbook per data source.

Configuration is read from an environment file. ``PURVIEW_ACCOUNT_NAME`` is
required. Service-principal authentication also requires
``AZURE_TENANT_ID``, ``AZURE_CLIENT_ID``, and ``AZURE_CLIENT_SECRET``.
"""

import argparse
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from azure.core.exceptions import AzureError
from azure.identity import ClientSecretCredential, DefaultAzureCredential
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


API_VERSION = "2023-09-01"
DEFAULT_ENV_FILE = "purview.env"
DEFAULT_OUTPUT_DIRECTORY = "reports"
DEFAULT_FILENAME_SUFFIX = "classifications"
AUTHENTICATION_MODES = ("azure-credential", "service-principal")
MODIFIED_TIME_RANGES = {
    "24h": "LAST_24H",
    "7d": "LAST_7D",
    "30d": "LAST_30D",
}
REQUEST_TIMEOUT_SECONDS = 60
ENTITY_GUID_BATCH_SIZE = 25
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SELECTED_FILL = PatternFill("solid", fgColor="D9EAD3")
HEADER_FONT = Font(name="Arial", color="FFFFFF", bold=True)
BODY_FONT = Font(name="Arial", color="000000")


class PurviewApiError(RuntimeError):
    """Report an expected Purview request, configuration, or output error."""

    pass


def parse_args():
    """Parse and validate command-line report options.

    Returns:
        argparse.Namespace: Validated command-line arguments.

    Notes:
        ``argparse`` terminates the process for invalid combinations. A custom
        qualified-name prefix is valid only for one named data source, while
        per-source output requires the data-source selection mode.
    """

    parser = argparse.ArgumentParser(
        description=(
            "Create an XLSX report of Microsoft Purview classifications by "
            "registered data source."
        )
    )
    source_selection = parser.add_mutually_exclusive_group(required=True)
    source_selection.add_argument(
        "--data-source",
        help=(
            "Registered Purview data source name to include in the report, or ALL "
            "to include every registered data source."
        ),
    )
    source_selection.add_argument(
        "--list-data-sources",
        action="store_true",
        help="Enumerate registered Purview data sources on screen and exit.",
    )
    parser.add_argument(
        "--filename-suffix",
        default=DEFAULT_FILENAME_SUFFIX,
        help=(
            "Filename suffix placed before the .xlsx extension "
            f"(default: {DEFAULT_FILENAME_SUFFIX})."
        ),
    )
    parser.add_argument(
        "--output-directory",
        default=DEFAULT_OUTPUT_DIRECTORY,
        help=(
            "Directory where generated XLSX files are stored "
            f"(default: {DEFAULT_OUTPUT_DIRECTORY})."
        ),
    )
    parser.add_argument(
        "--file-per-data-source",
        action="store_true",
        help=(
            "Generate a separate XLSX file for each data source in scope. Each "
            "filename is prefixed with the related data source name."
        ),
    )
    parser.add_argument(
        "--env-file",
        default=DEFAULT_ENV_FILE,
        help=f"Environment file containing PURVIEW_ACCOUNT_NAME (default: {DEFAULT_ENV_FILE}).",
    )
    parser.add_argument(
        "--authentication-mode",
        choices=AUTHENTICATION_MODES,
        default="azure-credential",
        help=(
            "Authentication method: azure-credential uses DefaultAzureCredential "
            "(Azure CLI, managed identity, developer tools, and other supported "
            "credentials); service-principal uses AZURE_TENANT_ID, AZURE_CLIENT_ID, "
            "and AZURE_CLIENT_SECRET (default: azure-credential)."
        ),
    )
    parser.add_argument(
        "--qualified-name-prefix",
        help=(
            "Optional qualifiedName prefix used to identify assets for the selected "
            "data source. Use this when the registration does not expose a usable endpoint."
        ),
    )
    parser.add_argument(
        "--page-size",
        type=int,
        default=1000,
        choices=range(1, 1001),
        metavar="[1-1000]",
        help="Purview discovery results per request (default: 1000).",
    )
    parser.add_argument(
        "--modified-within",
        choices=MODIFIED_TIME_RANGES,
        metavar="{24h,7d,30d}",
        help=(
            "Only include catalog assets modified within the previous 24 hours, "
            "7 days, or 30 days."
        ),
    )
    args = parser.parse_args()
    if args.qualified_name_prefix and (
        not args.data_source or args.data_source.casefold() == "all"
    ):
        parser.error(
            "--qualified-name-prefix requires a specific --data-source value."
        )
    if args.file_per_data_source and not args.data_source:
        parser.error("--file-per-data-source requires --data-source.")
    return args


def create_credential(authentication_mode):
    """Create the Azure credential selected for Purview API access.

    Args:
        authentication_mode (str): One of ``AUTHENTICATION_MODES``.

    Returns:
        azure.core.credentials.TokenCredential: Configured Azure credential.

    Raises:
        PurviewApiError: If service-principal configuration is incomplete or
            the authentication mode is unsupported.
    """

    if authentication_mode == "azure-credential":
        return DefaultAzureCredential()
    if authentication_mode != "service-principal":
        raise PurviewApiError(
            f"Unsupported authentication mode: {authentication_mode}"
        )

    variable_names = (
        "AZURE_TENANT_ID",
        "AZURE_CLIENT_ID",
        "AZURE_CLIENT_SECRET",
    )
    values = {name: os.getenv(name) for name in variable_names}
    missing = [name for name, value in values.items() if not value]
    if missing:
        raise PurviewApiError(
            "Service-principal authentication requires these environment "
            f"variables: {', '.join(missing)}."
        )

    return ClientSecretCredential(
        tenant_id=values["AZURE_TENANT_ID"],
        client_id=values["AZURE_CLIENT_ID"],
        client_secret=values["AZURE_CLIENT_SECRET"],
    )


def request_json(session, method, url, headers, **kwargs):
    """Send an HTTP request and decode its JSON response.

    Args:
        session (requests.Session): Session used to reuse HTTP connections.
        method (str): HTTP method, such as ``GET`` or ``POST``.
        url (str): Fully qualified Purview API URL.
        headers (dict): Request headers, including the bearer token.
        **kwargs: Additional arguments forwarded to ``Session.request``.

    Returns:
        Any: The JSON-compatible response body.

    Raises:
        PurviewApiError: If the request fails, returns an unsuccessful status,
            or contains invalid JSON.
    """

    try:
        response = session.request(
            method,
            url,
            headers=headers,
            timeout=REQUEST_TIMEOUT_SECONDS,
            **kwargs,
        )
        response.raise_for_status()
        return response.json()
    except requests.RequestException as exc:
        details = ""
        if exc.response is not None:
            details = f" Response: {exc.response.text[:1000]}"
        raise PurviewApiError(f"{method} {url} failed: {exc}.{details}") from exc
    except ValueError as exc:
        raise PurviewApiError(f"{method} {url} returned invalid JSON.") from exc


def list_data_sources(session, endpoint, headers):
    """Retrieve every registered data source from the scanning data plane.

    Args:
        session (requests.Session): Authenticated HTTP session.
        endpoint (str): Purview account endpoint without a trailing path.
        headers (dict): Request headers accepted by the Purview API.

    Returns:
        list[dict]: Data-source registration objects across all API pages.
    """

    url = f"{endpoint}/scan/datasources?api-version={API_VERSION}"
    data_sources = []

    while url:
        response = request_json(session, "GET", url, headers)
        data_sources.extend(response.get("value", []))
        next_link = response.get("nextLink")
        url = urljoin(endpoint, next_link) if next_link else None

    return data_sources


def list_catalog_assets(
    session,
    endpoint,
    headers,
    page_size,
    modified_within=None,
):
    """Yield catalog assets from the Purview discovery search API.

    Args:
        session (requests.Session): Authenticated HTTP session.
        endpoint (str): Purview account endpoint without a trailing path.
        headers (dict): Request headers accepted by the Purview API.
        page_size (int): Maximum number of assets requested per API page.
        modified_within (str | None): Optional CLI time-range key from
            ``MODIFIED_TIME_RANGES``.

    Yields:
        dict: One catalog asset from each paginated search response.

    Notes:
        The modified-time filter applies to the asset, not to the time at which
        a classification was assigned.
    """

    url = f"{endpoint}/datamap/api/search/query?api-version={API_VERSION}"
    continuation_token = None

    while True:
        body = {"keywords": None, "limit": page_size}
        if modified_within:
            body["filter"] = {
                "attributeName": "modifiedTime",
                "operator": "timerange",
                "attributeValue": MODIFIED_TIME_RANGES[modified_within],
            }
        if continuation_token:
            body["continuationToken"] = continuation_token

        response = request_json(session, "POST", url, headers, json=body)
        yield from response.get("value", [])

        continuation_token = response.get("continuationToken")
        if not continuation_token:
            break


def atlas_entity_to_asset(entity):
    """Convert an Atlas entity response into the report's asset shape.

    Args:
        entity (dict): Entity from ``entities`` or ``referredEntities``.

    Returns:
        dict: Normalized asset fields used by report generation.
    """

    attributes = entity.get("attributes") or {}
    relationship_attributes = entity.get("relationshipAttributes") or {}
    parent = {}
    for relationship_name in (
        "table",
        "view",
        "composeSchema",
        "tabular_schema",
    ):
        relationship = relationship_attributes.get(relationship_name)
        if isinstance(relationship, dict):
            parent = relationship
            break

    type_name = entity.get("typeName", "")
    return {
        "name": (
            attributes.get("name")
            or entity.get("displayText")
            or attributes.get("displayName")
            or ""
        ),
        "id": entity.get("guid", ""),
        "entityType": type_name,
        "assetType": [type_name] if type_name else [],
        "qualifiedName": attributes.get("qualifiedName", ""),
        "classification": entity.get("classifications") or [],
        "description": (
            attributes.get("description")
            or attributes.get("userDescription")
            or attributes.get("comment")
            or ""
        ),
        "collectionId": entity.get("collectionId", ""),
        "dataType": attributes.get("dataType") or attributes.get("type") or "",
        "parentName": parent.get("displayText", ""),
        "parentGuid": parent.get("guid", ""),
    }


def list_classified_columns(
    session,
    endpoint,
    headers,
    assets,
    batch_size=ENTITY_GUID_BATCH_SIZE,
):
    """Yield classified column entities from Atlas entity details.

    Discovery search records do not reliably contain child columns. This
    function retrieves matched entities in bulk and inspects both the requested
    ``entities`` and their expanded ``referredEntities`` for classified
    columns.

    Args:
        session (requests.Session): Authenticated HTTP session.
        endpoint (str): Purview account endpoint without a trailing path.
        headers (dict): Request headers accepted by the Purview API.
        assets (Iterable[dict]): Matched discovery assets whose relationships
            should be expanded.
        batch_size (int): Maximum GUIDs requested per bulk API call.

    Yields:
        dict: Normalized classified column asset.
    """

    guids = sorted(
        {
            str(asset.get("id"))
            for asset in assets
            if asset.get("id") not in (None, "")
        }
    )
    seen_columns = set()
    url = f"{endpoint}/datamap/api/atlas/v2/entity/bulk"

    for offset in range(0, len(guids), batch_size):
        batch = guids[offset : offset + batch_size]
        params = [
            ("api-version", API_VERSION),
            ("minExtInfo", "false"),
            ("ignoreRelationships", "false"),
        ]
        params.extend(("guid", guid) for guid in batch)
        response = request_json(
            session,
            "GET",
            url,
            headers,
            params=params,
        )
        entities = list(response.get("entities") or [])
        entities.extend((response.get("referredEntities") or {}).values())

        for entity in entities:
            asset = atlas_entity_to_asset(entity)
            if not is_column_asset(asset) or not classification_names(asset):
                continue
            identity = asset.get("id") or asset.get("qualifiedName")
            if not identity or identity in seen_columns:
                continue
            seen_columns.add(identity)
            yield asset


def normalize_locator(value):
    """Normalize an endpoint or qualified name for locator comparison.

    Normalization is case-insensitive, converts backslashes to slashes, removes
    URI schemes and query/fragment components, and strips outer slashes.

    Args:
        value (Any): Locator-like value; falsey values produce an empty string.

    Returns:
        str: A normalized host-and-path locator.
    """

    text = str(value or "").strip().lower().replace("\\", "/")
    if not text:
        return ""

    parsed = urlparse(text if "://" in text else f"//{text}")
    if parsed.netloc:
        text = f"{parsed.netloc}{parsed.path}"

    return text.strip("/")


def iter_nested_values(value, key_path=()):
    """Recursively yield string leaves and their key paths.

    Args:
        value (Any): Nested dictionary, list, or scalar to inspect.
        key_path (tuple[str, ...]): Path accumulated by recursive calls.

    Yields:
        tuple[tuple[str, ...], str]: The dictionary-key path and string value
        for each string leaf. List indexes are intentionally omitted.
    """

    if isinstance(value, dict):
        for key, child in value.items():
            yield from iter_nested_values(child, key_path + (str(key),))
    elif isinstance(value, list):
        for child in value:
            yield from iter_nested_values(child, key_path)
    elif isinstance(value, str):
        yield key_path, value


def data_source_locators(data_source):
    """Extract candidate asset locators from a data-source registration.

    Known endpoint-like properties are searched recursively. For services that
    register a child database, resource, or warehouse separately, that child
    name is also appended to the primary endpoint to create a more specific
    locator.

    Args:
        data_source (dict): Purview scanning data-source registration.

    Returns:
        list[str]: Unique normalized locators, longest first.
    """

    locator_keys = {
        "endpoint",
        "host",
        "hostname",
        "qualifiedname",
        "resourceuri",
        "serverendpoint",
        "url",
    }
    locators = set()

    for key_path, value in iter_nested_values(data_source.get("properties", {})):
        if key_path and key_path[-1].lower() in locator_keys:
            locator = normalize_locator(value)
            if locator:
                locators.add(locator)

    properties = data_source.get("properties") or {}
    endpoint = source_property(
        data_source,
        "endpoint",
        "serverEndpoint",
        "resourceUri",
        "url",
    )
    for child_name_key in ("databaseName", "resourceName", "warehouseName"):
        child_name = properties.get(child_name_key)
        if endpoint and isinstance(child_name, str) and child_name.strip():
            locators.add(
                normalize_locator(
                    f"{normalize_locator(endpoint)}/{child_name.strip('/')}"
                )
            )

    return sorted(locators, key=len, reverse=True)


def locator_matches(qualified_name, locator):
    """Return whether a locator identifies a qualified asset name.

    A match may be exact, a path prefix, or a complete path segment within the
    qualified name. Segment delimiters prevent partial host or path matches.

    Args:
        qualified_name (Any): Catalog asset qualified name.
        locator (Any): Registration locator to test.

    Returns:
        bool: ``True`` when the normalized locator identifies the asset.
    """

    qualified_name = normalize_locator(qualified_name)
    locator = normalize_locator(locator)
    if not qualified_name or not locator:
        return False

    return (
        qualified_name == locator
        or qualified_name.startswith(f"{locator}/")
        or f"/{locator}/" in f"/{qualified_name}/"
    )


def build_source_index(data_sources, selected_name=None, selected_prefix=None):
    """Build the ordered registration-to-locator index used for matching.

    Args:
        data_sources (list[dict]): All registered Purview data sources.
        selected_name (str | None): Optional selected registration name.
        selected_prefix (str | None): Optional locator override for the selected
            registration. It is preferred but does not discard discovered
            locators.

    Returns:
        list[tuple[dict, list[str]]]: Registrations paired with their locators.
        When a source is selected, its entry is sorted first.
    """

    source_index = []
    for source in data_sources:
        locators = data_source_locators(source)
        if (
            selected_name
            and source.get("name", "").casefold() == selected_name.casefold()
            and selected_prefix
        ):
            normalized_prefix = normalize_locator(selected_prefix)
            if normalized_prefix:
                locators = [normalized_prefix] + [
                    locator for locator in locators if locator != normalized_prefix
                ]
        source_index.append((source, locators))
    return sorted(
        source_index,
        key=lambda item: (
            bool(selected_name)
            and item[0].get("name", "").casefold() != selected_name.casefold()
        ),
    )


def match_asset_to_source(asset, source_index):
    """Find the most specific registered data source for an asset.

    Args:
        asset (dict): Purview catalog asset containing ``qualifiedName``.
        source_index (list[tuple[dict, list[str]]]): Registration locator index.

    Returns:
        dict | None: Registration with the longest matching locator, or ``None``
        when no locator matches.
    """

    qualified_name = asset.get("qualifiedName")
    matches = []

    for source, locators in source_index:
        for locator in locators:
            if locator_matches(qualified_name, locator):
                matches.append((len(locator), source))
                break

    if not matches:
        return None

    matches.sort(key=lambda item: item[0], reverse=True)
    return matches[0][1]


def classification_names(asset):
    """Return the unique classification names assigned to an asset.

    Purview responses may represent classifications as strings or objects.
    Object values prefer ``typeName`` and fall back to ``name``.

    Args:
        asset (dict): Catalog asset returned by discovery search.

    Returns:
        list[str]: Case-insensitively sorted, unique classification names.
    """

    names = set()
    for field_name in ("classification", "classifications"):
        classifications = asset.get(field_name) or []
        if not isinstance(classifications, (list, tuple, set)):
            classifications = [classifications]
        for classification in classifications:
            if isinstance(classification, str):
                name = classification
            elif isinstance(classification, dict):
                name = classification.get("typeName") or classification.get("name")
            else:
                name = None
            if name:
                names.add(name)
    return sorted(names, key=str.casefold)


def asset_type_names(asset):
    """Return the searchable entity and asset type names for an asset.

    Args:
        asset (dict): Catalog asset returned by discovery search.

    Returns:
        list[str]: Nonempty type names represented as strings.
    """

    values = [asset.get("entityType")]
    asset_types = asset.get("assetType") or []
    if isinstance(asset_types, list):
        values.extend(asset_types)
    else:
        values.append(asset_types)
    return [str(value) for value in values if value not in (None, "")]


def is_column_asset(asset):
    """Return whether Purview identifies an asset as a table-like column.

    Purview type names vary by connector (for example ``azure_sql_column`` or
    ``Database Column``), so entity and asset type names are matched on a
    complete ``column`` or ``columns`` word rather than a connector-specific
    allowlist.

    Args:
        asset (dict): Catalog asset returned by discovery search.

    Returns:
        bool: ``True`` when a type name contains a column type marker.
    """

    for name in asset_type_names(asset):
        normalized_name = re.sub(
            r"(?<=[a-z0-9])(?=[A-Z])",
            "_",
            name,
        ).casefold()
        if re.search(
            r"(^|[^a-z0-9])columns?($|[^a-z0-9])",
            normalized_name,
        ):
            return True
    return False


def formatted_asset_type(asset):
    """Return asset type values formatted for an Excel cell."""

    asset_type = asset.get("assetType") or []
    if isinstance(asset_type, list):
        return ", ".join(str(value) for value in asset_type)
    return str(asset_type)


def source_property(source, *keys):
    """Return the first populated registration property among candidate keys.

    Args:
        source (dict): Purview data-source registration.
        *keys (str): Property names in precedence order.

    Returns:
        Any: The first nonempty value. Dictionary values prefer
        ``referenceName`` and then ``name``; an empty string indicates no match.
    """

    properties = source.get("properties") or {}
    for key in keys:
        value = properties.get(key)
        if value not in (None, ""):
            if isinstance(value, dict):
                return (
                    value.get("referenceName")
                    or value.get("name")
                    or str(value)
                )
            return value
    return ""


def aggregate_assets(data_sources, assets, source_index):
    """Group assets by registration and calculate classification statistics.

    Args:
        data_sources (list[dict]): Registrations for which summaries are needed.
        assets (Iterable[dict]): Catalog assets to match and count.
        source_index (list[tuple[dict, list[str]]]): Registration locator index.

    Returns:
        tuple: ``(assets_by_source, summaries, unmatched_count)``, where the
        first value maps source names to matched assets, the second maps source
        names to count dictionaries, and the third counts unmatched assets.

    Notes:
        ``classification_assignments`` can exceed ``classified_assets`` because
        one asset may have multiple classifications. Unclassified assets are
        retained in summary counts but omitted from the detailed output sheet.
    """

    assets_by_source = defaultdict(list)
    unmatched_count = 0

    for asset in assets:
        source = match_asset_to_source(asset, source_index)
        if source is None:
            unmatched_count += 1
            continue
        assets_by_source[source.get("name", "")].append(asset)

    summaries = {}
    for source in data_sources:
        source_name = source.get("name", "")
        source_assets = assets_by_source[source_name]
        classification_counts = Counter()
        classified_assets = 0

        for asset in source_assets:
            names = classification_names(asset)
            if names:
                classified_assets += 1
                classification_counts.update(names)
            else:
                classification_counts["(Unclassified)"] += 1

        summaries[source_name] = {
            "matched_assets": len(source_assets),
            "classified_assets": classified_assets,
            "classification_assignments": sum(
                count
                for name, count in classification_counts.items()
                if name != "(Unclassified)"
            ),
            "classification_counts": classification_counts,
        }

    return assets_by_source, summaries, unmatched_count


def group_columns_by_source(columns, source_index):
    """Map normalized Atlas column entities to registered data sources."""

    columns_by_source = defaultdict(list)
    for column in columns:
        source = match_asset_to_source(column, source_index)
        if source is not None:
            columns_by_source[source.get("name", "")].append(column)
    return columns_by_source


def add_table(worksheet, name):
    """Convert a populated worksheet range into a styled Excel table.

    Args:
        worksheet (openpyxl.worksheet.worksheet.Worksheet): Target worksheet.
        name (str): Workbook-unique Excel table display name.

    Notes:
        Excel tables require at least one data row, so an empty worksheet gains
        a placeholder ``No records found`` row before the table is created.
    """

    if worksheet.max_row < 2:
        worksheet.append(["No records found"] + [""] * (worksheet.max_column - 1))

    table = Table(
        displayName=name,
        ref=f"A1:{worksheet.cell(worksheet.max_row, worksheet.max_column).coordinate}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def format_worksheet(worksheet, widths):
    """Apply shared header, body, view, and column-width formatting.

    Args:
        worksheet (openpyxl.worksheet.worksheet.Worksheet): Sheet to format.
        widths (dict[str, int | float]): Excel column letters mapped to widths.
    """

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 28

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def sanitize_filename_component(value):
    """Replace characters that are invalid in Windows filename components.

    Args:
        value (str): User-provided suffix or data-source name.

    Returns:
        str: Trimmed filename component with invalid characters replaced by
        underscores and trailing periods or spaces removed.
    """

    return re.sub(
        r'[<>:"/\\|?*\x00-\x1f]',
        "_",
        value.strip(),
    ).strip(" .")


def report_output_path(output_directory, filename_suffix):
    """Construct the base XLSX output path from CLI values.

    Args:
        output_directory (str | os.PathLike): Destination directory.
        filename_suffix (str): Requested workbook filename without extension.

    Returns:
        pathlib.Path: Sanitized ``.xlsx`` output path.

    Raises:
        PurviewApiError: If sanitization removes the entire filename suffix.
    """

    safe_filename_suffix = sanitize_filename_component(filename_suffix)
    if not safe_filename_suffix:
        raise PurviewApiError(
            "The filename suffix must contain at least one valid filename character."
        )
    return Path(output_directory) / f"{safe_filename_suffix}.xlsx"


def data_source_output_path(output_path, source_name):
    """Prefix a base output filename with a sanitized data-source name.

    Args:
        output_path (str | os.PathLike): Base combined-report output path.
        source_name (str): Registration name used as the filename prefix.

    Returns:
        pathlib.Path: Per-source path in the base path's directory.
    """

    output_path = Path(output_path)
    safe_source_name = sanitize_filename_component(source_name)
    if not safe_source_name:
        safe_source_name = "data-source"
    return output_path.with_name(f"{safe_source_name}-{output_path.name}")


def unique_output_path(output_path, used_output_paths):
    """Reserve a case-insensitively unique output path for this run.

    Args:
        output_path (pathlib.Path): Preferred workbook path.
        used_output_paths (set[str]): Absolute normalized paths already reserved.

    Returns:
        pathlib.Path: Preferred path or one suffixed with ``-2``, ``-3``, etc.

    Notes:
        This resolves collisions among generated names only. Existing files on
        disk are intentionally overwritten by ``Workbook.save``.
    """

    candidate = output_path
    suffix_number = 2
    normalized_path = str(candidate.absolute()).casefold()
    while normalized_path in used_output_paths:
        candidate = output_path.with_name(
            f"{output_path.stem}-{suffix_number}{output_path.suffix}"
        )
        suffix_number += 1
        normalized_path = str(candidate.absolute()).casefold()
    used_output_paths.add(normalized_path)
    return candidate


def create_workbook(
    output_path,
    report_sources,
    report_scope,
    source_index,
    assets_by_source,
    columns_by_source,
    summaries,
):
    """Create and save a formatted Purview classification workbook.

    Args:
        output_path (str | os.PathLike): Destination XLSX path.
        report_sources (list[dict]): Registrations included in this workbook.
        report_scope (str): Human-readable scope stored in workbook metadata.
        source_index (list[tuple[dict, list[str]]]): Registration locator index.
        assets_by_source (Mapping[str, list[dict]]): Matched assets by source.
        columns_by_source (Mapping[str, list[dict]]): Classified Atlas columns
            by source.
        summaries (Mapping[str, dict]): Aggregate counts by source.

    The workbook contains four sheets: registration metadata and totals,
    classification counts by source, one detail row per classified
    asset/classification pair, and one detail row per classified
    column/classification pair. Parent directories are created as needed.
    """

    workbook = Workbook()
    data_sources_sheet = workbook.active
    data_sources_sheet.title = "Data Sources"
    summary_sheet = workbook.create_sheet("Classification Summary")
    assets_sheet = workbook.create_sheet("Assets & Classifications")
    columns_sheet = workbook.create_sheet("Column Classifications")

    data_sources_sheet.append(
        [
            "Selected",
            "Data Source",
            "Kind",
            "ID",
            "Collection",
            "Location",
            "Resource Group",
            "Subscription ID",
            "Resource Name",
            "Endpoint",
            "Asset Match Locators",
            "Matched Assets",
            "Classified Assets",
            "Classification Assignments",
        ]
    )
    locator_map = {
        source.get("name", ""): locators for source, locators in source_index
    }
    for source in sorted(
        report_sources, key=lambda item: item.get("name", "").casefold()
    ):
        source_name = source.get("name", "")
        summary = summaries[source_name]
        data_sources_sheet.append(
            [
                "Yes",
                source_name,
                source.get("kind", ""),
                source.get("id", ""),
                source_property(source, "collection"),
                source_property(source, "location"),
                source_property(source, "resourceGroup"),
                source_property(source, "subscriptionId"),
                source_property(source, "resourceName"),
                source_property(
                    source,
                    "endpoint",
                    "serverEndpoint",
                    "resourceUri",
                    "url",
                ),
                "\n".join(locator_map[source_name]),
                summary["matched_assets"],
                summary["classified_assets"],
                summary["classification_assignments"],
            ]
        )
        for cell in data_sources_sheet[data_sources_sheet.max_row]:
            cell.fill = SELECTED_FILL

    summary_sheet.append(
        [
            "Data Source",
            "Kind",
            "Classification",
            "Asset Count",
        ]
    )
    source_by_name = {
        source.get("name", ""): source for source in report_sources
    }
    for source_name in sorted(source_by_name, key=str.casefold):
        counts = summaries[source_name]["classification_counts"]
        if not counts:
            summary_sheet.append(
                [
                    source_name,
                    source_by_name[source_name].get("kind", ""),
                    "(No matched assets)",
                    0,
                ]
            )
            continue
        for classification, count in sorted(
            counts.items(), key=lambda item: item[0].casefold()
        ):
            summary_sheet.append(
                [
                    source_name,
                    source_by_name[source_name].get("kind", ""),
                    classification,
                    count,
                ]
            )

    assets_sheet.append(
        [
            "Data Source",
            "Asset Name",
            "Asset GUID",
            "Entity Type",
            "Asset Type",
            "Qualified Name",
            "Classification",
            "Description",
            "Collection ID",
        ]
    )
    for source in sorted(
        report_sources, key=lambda item: item.get("name", "").casefold()
    ):
        source_name = source.get("name", "")
        source_assets = sorted(
            assets_by_source[source_name],
            key=lambda asset: (
                str(asset.get("qualifiedName", "")).casefold(),
                str(asset.get("name", "")).casefold(),
            ),
        )
        for asset in source_assets:
            classifications = classification_names(asset)
            if not classifications:
                continue
            for classification in classifications:
                assets_sheet.append(
                    [
                        source_name,
                        asset.get("name", ""),
                        asset.get("id", ""),
                        asset.get("entityType", ""),
                        formatted_asset_type(asset),
                        asset.get("qualifiedName", ""),
                        classification,
                        asset.get("description", ""),
                        asset.get("collectionId", ""),
                    ]
                )

    columns_sheet.append(
        [
            "Data Source",
            "Column Name",
            "Column GUID",
            "Entity Type",
            "Asset Type",
            "Data Type",
            "Parent Asset",
            "Parent GUID",
            "Column Qualified Name",
            "Classification",
            "Description",
            "Collection ID",
        ]
    )
    for source in sorted(
        report_sources, key=lambda item: item.get("name", "").casefold()
    ):
        source_name = source.get("name", "")
        column_assets = sorted(
            columns_by_source[source_name],
            key=lambda asset: (
                str(asset.get("qualifiedName", "")).casefold(),
                str(asset.get("name", "")).casefold(),
            ),
        )
        for column in column_assets:
            for classification in classification_names(column):
                columns_sheet.append(
                    [
                        source_name,
                        column.get("name", ""),
                        column.get("id", ""),
                        column.get("entityType", ""),
                        formatted_asset_type(column),
                        column.get("dataType", ""),
                        column.get("parentName", ""),
                        column.get("parentGuid", ""),
                        column.get("qualifiedName", ""),
                        classification,
                        column.get("description", ""),
                        column.get("collectionId", ""),
                    ]
                )

    add_table(data_sources_sheet, "DataSourcesTable")
    add_table(summary_sheet, "ClassificationSummaryTable")
    add_table(assets_sheet, "AssetsClassificationsTable")
    add_table(columns_sheet, "ColumnClassificationsTable")
    format_worksheet(
        data_sources_sheet,
        {
            "A": 10,
            "B": 28,
            "C": 25,
            "D": 42,
            "E": 24,
            "F": 18,
            "G": 22,
            "H": 38,
            "I": 28,
            "J": 45,
            "K": 55,
            "L": 16,
            "M": 18,
            "N": 25,
        },
    )
    format_worksheet(
        summary_sheet,
        {"A": 30, "B": 28, "C": 45, "D": 16},
    )
    format_worksheet(
        assets_sheet,
        {
            "A": 28,
            "B": 35,
            "C": 38,
            "D": 30,
            "E": 30,
            "F": 20,
            "G": 35,
            "H": 38,
            "I": 70,
            "J": 45,
            "K": 55,
            "L": 38,
        },
    )
    format_worksheet(
        columns_sheet,
        {
            "A": 28,
            "B": 35,
            "C": 38,
            "D": 30,
            "E": 30,
            "F": 70,
            "G": 45,
            "H": 55,
            "I": 38,
        },
    )

    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    workbook.properties.title = "Microsoft Purview data source classifications"
    workbook.properties.subject = (
        f"Classification inventory for data source scope {report_scope}"
    )
    workbook.properties.creator = "PurviewDataSourceClassifications.py"
    workbook.properties.description = (
        f"Generated {generated_at}. Assets are mapped to registered data sources "
        "using the longest endpoint or qualified-name locator from registration metadata."
    )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main():
    """Run authentication, collection, aggregation, and report generation.

    Returns:
        int: Process exit code: ``0`` for success and ``1`` for an expected
        Azure, Purview, filesystem, or configuration failure.

    Side Effects:
        Loads environment variables, calls Microsoft Purview APIs, prints
        progress and warnings, creates output directories, and writes XLSX
        files. ``--list-data-sources`` prints registrations and exits before
        catalog retrieval.
    """

    args = parse_args()
    load_dotenv(args.env_file)
    account_name = os.getenv("PURVIEW_ACCOUNT_NAME")
    if not account_name:
        print(
            f"PURVIEW_ACCOUNT_NAME was not found in {args.env_file} or the environment.",
            file=sys.stderr,
        )
        return 1

    endpoint = f"https://{account_name}.purview.azure.com"
    try:
        credential = create_credential(args.authentication_mode)
        access_token = credential.get_token("https://purview.azure.net/.default").token
        headers = {
            "Authorization": f"Bearer {access_token}",
            "Content-Type": "application/json",
        }

        with requests.Session() as session:
            print("Retrieving registered data sources...")
            data_sources = list_data_sources(session, endpoint, headers)
            if args.list_data_sources:
                sorted_sources = sorted(
                    data_sources,
                    key=lambda source: source.get("name", "").casefold(),
                )
                if not sorted_sources:
                    print("No registered data sources found.")
                for number, source in enumerate(sorted_sources, start=1):
                    kind = source.get("kind", "")
                    kind_suffix = f" ({kind})" if kind else ""
                    print(f"{number}. {source.get('name', '')}{kind_suffix}")
                return 0

            include_all_sources = args.data_source.casefold() == "all"
            selected_source = None
            if not include_all_sources:
                selected_source = next(
                    (
                        source
                        for source in data_sources
                        if source.get("name", "").casefold()
                        == args.data_source.casefold()
                    ),
                    None,
                )
            if not include_all_sources and selected_source is None:
                available = ", ".join(
                    sorted(
                        (source.get("name", "") for source in data_sources),
                        key=str.casefold,
                    )
                )
                raise PurviewApiError(
                    f"Data source '{args.data_source}' was not found. "
                    f"Available data sources: {available or '(none)'}"
                )

            report_sources = (
                data_sources if include_all_sources else [selected_source]
            )
            source_index = build_source_index(
                data_sources,
                None if include_all_sources else selected_source.get("name", ""),
                args.qualified_name_prefix,
            )
            if selected_source is not None:
                selected_locators = next(
                    locators
                    for source, locators in source_index
                    if source is selected_source
                )
                if not selected_locators:
                    raise PurviewApiError(
                        "The selected registration has no endpoint metadata that can be "
                        "matched to asset qualified names. Rerun with "
                        "--qualified-name-prefix."
                    )

            print("Retrieving catalog assets...")
            assets = list(
                list_catalog_assets(
                    session,
                    endpoint,
                    headers,
                    args.page_size,
                    args.modified_within,
                )
            )
            assets_by_source, summaries, unmatched_count = aggregate_assets(
                data_sources,
                assets,
                source_index,
            )
            report_source_names = {
                source.get("name", "") for source in report_sources
            }
            detail_assets = [
                asset
                for source_name in report_source_names
                for asset in assets_by_source[source_name]
            ]
            print("Retrieving column classifications...")
            classified_columns = list(
                list_classified_columns(
                    session,
                    endpoint,
                    headers,
                    detail_assets,
                )
            )
            columns_by_source = group_columns_by_source(
                classified_columns,
                source_index,
            )

        base_output_path = report_output_path(
            args.output_directory,
            args.filename_suffix,
        )
        created_reports = []
        if args.file_per_data_source:
            used_output_paths = set()
            for source in sorted(
                report_sources,
                key=lambda item: item.get("name", "").casefold(),
            ):
                source_name = source.get("name", "")
                output_path = unique_output_path(
                    data_source_output_path(base_output_path, source_name),
                    used_output_paths,
                )
                create_workbook(
                    output_path,
                    [source],
                    source_name,
                    source_index,
                    assets_by_source,
                    columns_by_source,
                    summaries,
                )
                created_reports.append((source_name, output_path))
        else:
            create_workbook(
                base_output_path,
                report_sources,
                "ALL" if include_all_sources else selected_source.get("name", ""),
                source_index,
                assets_by_source,
                columns_by_source,
                summaries,
            )
            created_reports.append(
                (
                    "ALL"
                    if include_all_sources
                    else selected_source.get("name", ""),
                    base_output_path,
                )
            )

        report_asset_count = sum(
            len(assets_by_source[source.get("name", "")])
            for source in report_sources
        )
        report_scope = (
            "all registered data sources"
            if include_all_sources
            else f"'{selected_source.get('name', '')}'"
        )
        if args.file_per_data_source:
            print(
                f"Created {len(created_reports)} reports with {report_asset_count} "
                f"assets for {report_scope}:"
            )
            for source_name, output_path in created_reports:
                print(f"- {source_name}: {output_path.resolve()}")
        else:
            print(
                f"Created {created_reports[0][1].resolve()} with "
                f"{report_asset_count} assets for {report_scope}."
            )
        if unmatched_count:
            print(
                f"Warning: {unmatched_count} catalog assets could not be mapped to a "
                "registered data source using registration endpoint metadata.",
                file=sys.stderr,
            )
        return 0
    except (AzureError, PurviewApiError, OSError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
